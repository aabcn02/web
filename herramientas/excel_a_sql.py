"""Convierte el Excel de grupos del área (GRUPOS_PAGINA_WEB.xlsx) en cargar-grupos.sql.

Uso:  python3 herramientas/excel_a_sql.py ruta/al/excel.xlsx  > cargar-grupos.sql
Solo se usa para la carga inicial; después todo se edita en el panel.
"""
import re
import sys

import openpyxl


def txt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def sql(v):
    if v is None or v == "":
        return "NULL"
    if isinstance(v, (int, float)):
        return str(v)
    return "'" + str(v).replace("'", "''") + "'"


TEL_RE = re.compile(r"\s*TEL[.:]?\s*(.+)$", re.I)


def separar_telefono(obs):
    """'Sesiona ... TEL: 686 311 45 88' -> ('Sesiona ...', '686 311 45 88')"""
    m = TEL_RE.search(obs)
    if not m:
        return obs, ""
    numeros = re.split(r"\s+y\s+", m.group(1).strip().rstrip("."), flags=re.I)
    numeros = [re.sub(r"\s+", " ", n).strip() for n in numeros if re.search(r"\d{3}", n)]
    return obs[: m.start()].strip(), " / ".join(numeros)


def limpiar_horario(obs):
    h = re.sub(r"^Sesionas?\s+(de\s+)?", "", obs, flags=re.I)
    h = re.sub(r"\(\s*(invierno|verano)\s*\)", r"(\1)", h, flags=re.I)
    h = re.sub(r"\s+([.,;)])", r"\1", h)
    h = re.sub(r"(\d):\s+(\d\d)", r"\1:\2", h)          # "7: 00" -> "7:00"
    h = re.sub(r"(?<=[^\s.]) (Y|De)", lambda m: f" {m.group(1).lower()} ", h)
    h = re.sub(r"\s{2,}", " ", h).strip()
    return h[:1].upper() + h[1:] if h else h


DIAS = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]


def hora12(hhmm):
    h, m = (int(x) for x in hhmm.split(":"))
    suf = "am" if h < 12 else "pm"
    return f"{(h - 1) % 12 + 1}:{m:02d} {suf}"


def rango(t):
    """'07:30 - 09:00' -> '7:30 a 9:00 am'"""
    m = re.fullmatch(r"(\d{1,2}:\d{2})\s*-\s*(\d{1,2}:\d{2})", t)
    if not m:
        return t
    a, b = hora12(m.group(1)), hora12(m.group(2))
    if a[-2:] == b[-2:]:
        a = a[:-3]
    return f"{a} a {b}"


def horario_por_dias(r, i0):
    """Junta días seguidos con el mismo horario: 'Lunes a viernes de ...; sábado y domingo de ...'"""
    vals = [txt(r[i0 + k]) for k in range(7)]
    tramos, k = [], 0
    while k < 7:
        if not vals[k]:
            k += 1
            continue
        j = k
        while j + 1 < 7 and vals[j + 1] == vals[k]:
            j += 1
        dias = DIAS[k] if j == k else (f"{DIAS[k]} y {DIAS[j]}" if j == k + 1 else f"{DIAS[k]} a {DIAS[j]}")
        tramos.append(f"{dias.lower()} de {rango(vals[k])}")
        k = j + 1
    h = "; ".join(tramos)
    return h[:1].upper() + h[1:]


def main(ruta):
    ws = openpyxl.load_workbook(ruta, data_only=True).active
    enc = [txt(c.value) for c in ws[1]]
    col = {n: i for i, n in enumerate(enc)}
    filas = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        grupo = txt(r[col["Nombre Grupo"]])
        distrito = txt(r[col["Distrito"]])
        if not grupo or not distrito:  # salta renglones vacíos y "Informe generado..."
            continue
        obs = txt(r[col["Observación horarios"]])
        obs, tel = separar_telefono(obs)
        horario = limpiar_horario(obs)
        if not horario or not re.search(r"\d:\d", horario):
            por_dias = horario_por_dias(r, col["Horario Lunes"])
            horario = por_dias or horario
        ciudad = txt(r[col["Municipio"]])
        filas.append({
            "distrito": distrito,
            "grupo": grupo,
            "ciudad": ciudad,
            "estado": txt(r[col["Estado"]]),
            "direccion": txt(r[col["Domicilio"]]),
            "colonia": txt(r[col["Colonia"]]),
            "cp": txt(r[col["Código postal"]]),
            "referencia": txt(r[col["Referencia"]]),
            "terapia": txt(r[col["Tipo"]]),
            "personas": txt(r[col["Genero"]]),
            "idioma": "Inglés" if txt(r[col["Idioma"]]).lower().startswith("ingl") else txt(r[col["Idioma"]]),
            "horario": horario,
            "telefono": tel,
            "maps_url": txt(r[col["mapa"]]) if "mapa" in col else "",
        })

    campos = list(filas[0].keys())
    # Sin comentarios: la consola de D1 los rechaza
    esquema = "\n".join(l for l in open(__file__.replace("herramientas/excel_a_sql.py", "schema.sql"), encoding="utf-8").read().splitlines()
                       if not l.lstrip().startswith("--"))
    print(esquema.replace("CREATE TABLE IF NOT EXISTS grupos", "DROP TABLE IF EXISTS grupos;\nCREATE TABLE grupos"))
    print(f"INSERT INTO grupos ({', '.join(campos)}) VALUES")
    print(",\n".join("(" + ", ".join(sql(f[c]) for c in campos) + ")" for f in filas) + ";")
    print(f"-- {len(filas)} grupos", file=sys.stderr)


if __name__ == "__main__":
    main(sys.argv[1])
