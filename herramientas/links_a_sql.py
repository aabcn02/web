"""Lee la columna 'mapa' (links de Google Maps) del Excel y genera actualizar-ubicaciones.sql.

Uso:  python3 herramientas/links_a_sql.py excel.xlsx > actualizar-ubicaciones.sql
Solo cambia el link y borra el pin anterior de cada grupo, para que /api/ubicar lo vuelva a poner
con el link. No borra grupos ni toca ningún otro dato.
"""
import re
import sys

import openpyxl


def txt(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v or "")).strip()


def sql(v):
    return "'" + str(v).replace("'", "''") + "'"


ws = openpyxl.load_workbook(sys.argv[1]).active
enc = [txt(c.value) for c in ws[1]]
col = {n: i for i, n in enumerate(enc)}
n = 0
for fila in ws.iter_rows(min_row=2):
    grupo, distrito = txt(fila[col["Nombre Grupo"]].value), txt(fila[col["Distrito"]].value)
    celda = fila[col["mapa"]]
    link = txt((celda.hyperlink.target if celda.hyperlink else None) or celda.value)
    if not grupo or not distrito or not re.match(r"^https://", link):
        continue
    print(f"UPDATE grupos SET maps_url = {sql(link)}, lat = NULL, lng = NULL, precision_ubic = NULL, geo_intento = 0 "
          f"WHERE grupo = {sql(grupo)} AND distrito = {sql(distrito)};")
    n += 1
print(f"-- {n} links", file=sys.stderr)
